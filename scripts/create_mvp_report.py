"""Create the concise decision report using the bundled python-docx runtime."""
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parents[1]
DEST=ROOT/"deliverables/Live translation Bao cao MVP.docx"
doc=Document();section=doc.sections[0]
section.page_width=Inches(8.5);section.page_height=Inches(11)
section.top_margin=Inches(.7);section.bottom_margin=Inches(.7)
section.left_margin=Inches(.8);section.right_margin=Inches(.8)
for name in ("Normal","Title","Subtitle","Heading 1","Heading 2","List Bullet","List Number"):
    style=doc.styles[name];style.font.name="Calibri";style.font.color.rgb=RGBColor(0,0,0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"),"Yu Gothic")
normal=doc.styles["Normal"];normal.font.size=Pt(11)
normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.08
for name,size in (("Title",23),("Heading 1",15),("Heading 2",12)):
    doc.styles[name].font.size=Pt(size)
    doc.styles[name].paragraph_format.space_before=Pt(11)
    doc.styles[name].paragraph_format.space_after=Pt(6)
doc.styles["Title"].paragraph_format.space_before=Pt(0)
for style in doc.styles:
    for border in style.element.findall('.//'+qn("w:pBdr")):
        border.getparent().remove(border)
doc.core_properties.title="Quyết định và lộ trình MVP dịch manga Nhật Việt"
doc.core_properties.author="Live translation"
doc.core_properties.comments=""


def p(value,style=None):return doc.add_paragraph(value,style)
def h(value):doc.add_heading(value,level=1)
def link(paragraph,label,url):
    item=OxmlElement("w:hyperlink");item.set(qn("r:id"),paragraph.part.relate_to(url,RT.HYPERLINK,is_external=True))
    run=OxmlElement("w:r");pr=OxmlElement("w:rPr")
    color=OxmlElement("w:color");color.set(qn("w:val"),"24526B");pr.append(color)
    run.append(pr);text=OxmlElement("w:t");text.text=label;run.append(text);item.append(run);paragraph._p.append(item)


def table(headers,rows,widths):
    table=doc.add_table(rows=1,cols=len(headers));table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
    for column,width in zip(table.columns,widths):column.width=Inches(width)
    borders=OxmlElement("w:tblBorders")
    for side in ("top","left","bottom","right","insideH","insideV"):
        edge=OxmlElement("w:"+side)
        for key,value in (("val","single"),("sz","4"),("color","D9D9D9")):edge.set(qn("w:"+key),value)
        borders.append(edge)
    table._tbl.tblPr.append(borders)
    for index,values in enumerate([headers,*rows]):
        row=table.rows[0] if index==0 else table.add_row()
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement("w:cantSplit"))
        if index==0:trpr.append(OxmlElement("w:tblHeader"))
        for j,(cell,value) in enumerate(zip(row.cells,values)):
            cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text=value
            shade=OxmlElement("w:shd");shade.set(qn("w:fill"),"E6EEF2" if index==0 else ("F7F9FA" if index%2==0 else "FFFFFF"));cell._tc.get_or_add_tcPr().append(shade)
            margins=OxmlElement("w:tcMar")
            for side in ("top","left","bottom","right"):
                m=OxmlElement("w:"+side);m.set(qn("w:w"),"95");m.set(qn("w:type"),"dxa");margins.append(m)
            cell._tc.get_or_add_tcPr().append(margins)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after=Pt(0);paragraph.paragraph_format.line_spacing=1.05
                for run in paragraph.runs:run.font.size=Pt(10.5);run.bold=index==0
                if j==len(headers)-1 and len(headers)==3 and index>0:paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p("").paragraph_format.space_after=Pt(0)


doc.add_heading("Quyết định và lộ trình MVP dịch manga Nhật Việt",0)
p("Cập nhật ngày 13 tháng 10 năm 2026")
p("Chốt extension Chrome/Edge với companion Windows, OCR chạy local và dịch text bằng backend người đọc chọn. Giữ chất Nhật trong cách gọi, câu Việt tự nhiên và bộ nhớ riêng theo từng truyện. Bản development đã chạy xuyên suốt trong kiểm thử native riêng; chưa đủ bằng chứng để công bố chất lượng chuẩn dịch giả hoặc gần tức thì trên trang mới.")
h("Cấu hình chọn cho bản thử")
table(["Thành phần","Lựa chọn và điều kiện"],[
    ["Phát hiện và OCR","PP-OCRv5 mobile detector + manga-ocr. Chỉ đọc crop; ảnh ở RAM. Không coi đây là OCR tốt nhất trên mọi truyện."],
    ["Dịch offline","Hy-MT2 1.8B Q6_K qua llama.cpp CPU. Gọn để thử, nhưng cần tiếp tục đối chiếu các lỗi chủ thể, tình thái và văn phong."],
    ["Dịch cloud","DeepSeek V4.1 Flash qua OpenRouter là ứng viên ưu tiên từ pilot nhỏ; Gemini Flash-Lite là đối chứng. BYOK gửi text/context, không gửi ảnh mặc định."],
    ["Bộ nhớ nền","Qwen3 4B Instruct Q4_K_M. Tóm tắt tự động đang tắt mặc định vì đã gặp suy diễn sai; lời thoại gốc được giữ làm bằng chứng."],
],[1.35,5.55])
h("Luồng đọc ít thao tác")
p("Lần đầu: xác nhận truyện, chapter và giọng dịch; có thể bổ sung recap riêng. Sau đó bật dịch ảnh đang nhìn. Thư viện giữ tiến độ và hồ sơ, nên không phải nhập lại từ đầu. Tùy chọn chuẩn bị tối đa hai ảnh đã tải giúp đọc tiếp; không crawl chapter ẩn.")
p("Overlay bám phần tử ảnh, theo cuộn và resize; bóng khó dùng nhãn thay vì che tranh. Bấm bóng để xem gốc, sửa OCR hoặc người nói. Tên AI đề xuất có Đúng / Đổi / Để sau, không chặn đọc; tên đã xác nhận sửa lại được với mốc áp dụng.")
h("Nguyên tắc dịch")
p("Giữ hậu tố và cách gọi có trong nguồn như -kun, -san, senpai, onii-chan; không tự thêm. Tách họ, tên và biệt danh, không suy tình cảm từ việc đổi cách gọi. Xưng hô A→B chỉ áp dụng khi biết người nói và người nghe; chưa rõ thì lược đại từ tự nhiên. Hán–Việt dùng chọn lọc cho tên kỹ năng/vật phẩm fantasy; glossary có phiên bản thắng phong cách. Gia phả Việt phải có bằng chứng, không đoán cô/dì hay nội/ngoại.")

doc.add_page_break();h("Bộ nhớ và số đo thực tế")
p("SQLite là dữ liệu chuẩn; Markdown là bản đọc/export cho AI và người dùng. Mỗi fact, tên và sự kiện có nguồn cùng mốc chapter/trang. Hồ sơ truyện, xưng hô và context dài hạn giữ đến khi người đọc xóa; cache OCR/dịch mặc định 24 giờ, tổng hạn 100 MB. API key nằm trong kho bí mật Windows. Không lưu ảnh vào hồ sơ.")
p("Không gửi cả lịch sử 100 chapter mỗi request. Chỉ lấy nhân vật, thuật ngữ, hội thoại gần và bằng chứng liên quan trong ngân sách ngắn. Khi đọc lại chapter 20, sự kiện biết ở chapter 80 bị loại. Tóm tắt AI chưa xác nhận không trở thành sự kiện chắc chắn: prompt dùng lời thoại gốc có nguồn, không dùng câu diễn giải sai như đã về trước thay cho cứ về trước đi.")
p("Cache kiểm tra lại context thực sự dùng cho từng nhóm thoại; đổi tên liên quan làm mất hiệu lực bản dịch, ghi chú không liên quan không buộc dịch lại. Hồ sơ JSON riêng tư giữ thể loại, alias, tiến độ và lịch sử; chuyển theo phần, tối đa 16 MB. Nhập lại không tạo bản ghi trùng; giữ vị trí đang đọc và xác nhận đã khóa. Nhóm tóm tắt giữ đúng liên kết nguồn khi đổi ID; thiếu nguồn hoặc vượt mốc đọc thì bị từ chối. Thư viện có tìm kiếm và ghi chú dài mở khi cần.")
table(["Kiểm thử","Kết quả đã đo","Phạm vi"],[
    ["Native OCR và dịch warm","7,46 giây / trang 4 bóng","CPU 4 luồng"],
    ["Cache dịch local","3,02 ms","Chưa tính vẽ UI"],
    ["Hồ sơ 1,65 MB qua native","Nhập 183 ms / xuất 92 ms","14 phần"],
    ["Toàn trang 12 bóng trong pilot","Khoảng 24–33 giây","Không phải p95"],
    ["OCR crop có ground truth","CER 9,37%; đúng hết 61,54%","78 crop"],
],[2.15,2.8,1.95])
p("Các phép đo khác nhau không phải một leaderboard hay p95 đại diện. Hồ sơ 100 chapter chứa 1.000 tóm tắt tự tạo để thử lưu trữ, không phải kết quả AI đọc 100 chapter. Tập riêng chỉ có ba bản dịch Việt chuẩn. Số cache không chứng minh tốc độ trang mới; model nhỏ vẫn sai chủ thể, cấu trúc sai khiến và xưng hô trong regression ngữ nghĩa.")
h("Đã kiểm tra và còn phải đạt")
p("98 test Python và 11 test TypeScript đạt. QA overlay mô phỏng engine đạt ở 1280, 760 và 390 px: cuộn lồng nhau, object-fit, sửa tên/OCR, hủy kết quả cũ và prefetch giới hạn. QA thư viện kiểm tra tìm kiếm, mở ghi chú, đổi truyện khi phản hồi cũ về muộn và nhập UTF-8 lớn. Native thật xác nhận OCR, dịch, cache và chuyển hồ sơ có giới hạn.")
p("Edge cá nhân chưa xác minh overlay toàn luồng. Đăng ký host tồn tại trong tài khoản Windows thật, nhưng phiên trước popup báo không tìm thấy host. Lượt computer-use mới nhận ra cửa sổ test song không kích hoạt được; chưa thao tác dịch. QA giả lập không thay thế test live này. Bộ cài còn phụ thuộc .venv của máy phát triển, chưa portable. Tóm tắt chapter/nhóm chapter và chuỗi 100 chapter bằng model thật cần hoàn thiện, không chỉ kiểm tra dữ liệu nhập tay.")

doc.add_page_break();h("Bản quyền và phát hành")
p("Offline, BYOK, ảnh chỉ ở RAM và disclaimer không tạo miễn trừ bản quyền. Điều 4, 20 và 25 của nguồn [1] xác định quyền làm tác phẩm phái sinh và ngoại lệ có điều kiện, không cấp phép chung dịch toàn truyện. Trước beta công khai hoặc thu phí, đối chiếu bản hiện hành [4] và rà quyền nội dung, lưu context, điều khoản website, dữ liệu cá nhân, nghĩa vụ thương mại với người có chuyên môn pháp lý.")
p("Manga109 chỉ dùng học thuật phi thương mại. Manga109-s cho phép thương mại hóa kết quả ML theo điều kiện, không cho đóng gói/phân phối ảnh dataset hoặc ảnh sửa thành sản phẩm. Dùng đúng giấy phép cho benchmark nội bộ; demo công khai dùng nội dung tự tạo hoặc được cấp phép. Giấy phép model không cấp quyền đối với truyện đầu vào. [2]")
p("Điều khoản API DeepSeek cho phép tích hợp ứng dụng phục vụ người dùng bên ngoài; không đồng nghĩa quyền bán key/tài khoản hoặc relay API tùy ý. Gói 30 ngày credit hữu hạn chỉ mở sau beta và rà hợp đồng. Key doanh nghiệp ở backend, không ở extension; công bố dữ liệu gửi, retention, giới hạn gói và hoàn tiền. BYOK không chuyển toàn bộ trách nhiệm sang reader. [3]")
h("Thứ tự hoàn thiện")
for value in [
    "Chốt chất lượng trước: mở rộng tập Nhật–Việt có quyền và chấm mù lỗi chủ thể, phủ định, tình thái, danh xưng; benchmark OCR toàn trang thay vì chỉ crop chuẩn.",
    "Tối ưu tương tác: đo nhóm đầu 2/4, prompt và CPU từng biến; giữ model nóng, ưu tiên viewport, hủy prefetch khi không còn đọc. Không đánh đổi nghĩa để đạt số nhanh.",
    "Hoàn thiện bộ nhớ: tổng hợp chapter có trạng thái đọc dở/đầy đủ, nguồn bền vững, giới hạn spoiler; kiểm thử 100 chapter, sửa dữ kiện và cache phụ thuộc.",
    "Đóng gói rồi beta: companion portable, cập nhật có checksum/chữ ký, test live máy sạch và các kiểu ảnh/canvas. Không thao tác phần mềm bảo mật để vượt cảnh báo.",
    "Sau beta: cloud có credit và thanh toán; Android reader/share-in, iPhone reader hoặc Safari extension. Giao diện 390 px không chứng minh chạy model trên điện thoại.",
]:p(value,"List Number")
p("Chat nhắc nhân vật và giải thích sự kiện lưu trong roadmap, chưa triển khai. Khi làm phải dùng mốc đã đọc, phân biệt bằng chứng và suy đoán; mở spoiler chỉ sau cảnh báo và xác nhận riêng.")
h("Nguồn chính")
sources=[
    ("[1] Luật SHTT hợp nhất trong Công báo 10 04 2026","https://congbaocdn.chinhphu.vn/180507251028987904/2026/4/10/469197-1775632354_v1_1775783230_signed.pdf"),
    ("[2] Đại học Tokyo Manga109 License and Terms","https://manga109.github.io/manga109-project-website/en/"),
    ("[3] DeepSeek Open Platform Terms of Service","https://cdn.deepseek.com/policies/en-US/deepseek-open-platform-terms-of-service.html"),
    ("[4] Văn bản hợp nhất 160 2026 ngày 29 09 2026","https://congbao.chinhphu.vn/van-ban/van-ban-hop-nhat-so-160-2026-vbhn-lq-vpqh-470657/69077.htm"),
    ("manga-ocr công cụ OCR chuyên manga","https://github.com/kha-white/manga-ocr"),
    ("Japan Foundation Irodori bài 3 tên và cách gọi","https://www.irodori.jpf.go.jp/assets/data/starter/pdf/X_L03.pdf"),
]
for label,url in sources:
    paragraph=p("");paragraph.paragraph_format.space_after=Pt(3);link(paragraph,label,url)
    for run in paragraph.runs:run.font.size=Pt(10)
DEST.parent.mkdir(parents=True,exist_ok=True);doc.save(DEST)
print(str(DEST))
